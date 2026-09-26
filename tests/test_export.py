"""Tests for the T20 export bundle (xlsx + manifest + LLM chart spec).

The chart decision is stubbed at the ``export.chart_spec.chat`` boundary
(same convention as tests/test_qa.py stubbing ``nl2data.qa.generate``); the
xlsx roundtrip is verified by reading the workbook back with openpyxl.
"""

from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from exec.runner import ExecutionResult
from export import run_export
from export.chart_spec import (
    ChartSpec,
    decide_chart_spec,
    validate_spec,
)
from guard.validate import ValidatedSQL
from llm.chat import ChatResult, LlmError, Usage
from nl2data.cli import app
from nl2data.config import Nl2DataConfig
from nl2data.config import load_config as _load
from nl2data.qa import QaOutcome

runner = CliRunner()
FIXTURE_YAML = "paths:\n  data_dir: data\n"

_BAR_SPEC = {
    "chart_type": "bar",
    "dimension": "borough",
    "measures": ["avg_fare"],
    "title": "各行政区平均车费",
    "top_n": 15,
}


def _execution(
    columns: list[str],
    rows: list[dict[str, object]],
    rowcount: int | None = None,
    detail_ref: str | None = None,
) -> ExecutionResult:
    return ExecutionResult(
        status="ok",
        rowcount=rowcount if rowcount is not None else len(rows),
        columns=columns,
        rows=rows,
        profile=None,
        detail_ref=detail_ref,
        error=None,
        latency_ms=7.0,
    )


def _outcome(execution: ExecutionResult, sql: str = "SELECT 1") -> QaOutcome:
    return QaOutcome(
        question="q",
        retrieved_tables=["t"],
        vsql=ValidatedSQL(sql=sql, tables=["t"], limit=500),
        execution=execution,
    )


def _result(parsed: dict | None, content: str = "{}") -> ChatResult:
    return ChatResult(
        content=content, parsed=parsed, usage=Usage(0, 0, 0), model="stub", latency_ms=1.0
    )


def _install_chat(monkeypatch: pytest.MonkeyPatch, chat_result: ChatResult) -> list[dict]:
    """Stub export.chart_spec.chat; returns the captured messages list."""
    calls: list[dict] = []

    def _fake(messages, json_schema=None, cfg=None):  # noqa: ANN001, ANN202
        calls.append({"messages": messages, "json_schema": json_schema})
        return chat_result

    monkeypatch.setattr("export.chart_spec.chat", _fake)
    return calls


@pytest.fixture()
def cfg(tmp_path: Path) -> Nl2DataConfig:
    """Isolated config rooted at tmp_path (export dir defaults under it)."""
    config_file = tmp_path / "config.yaml"
    config_file.write_text(FIXTURE_YAML, encoding="utf-8")
    loaded = _load(config_file)
    assert loaded.export.dir.is_absolute() and str(tmp_path) in str(loaded.export.dir)
    return loaded


# ------------------------------------------------------------- validator ----
def _exec_for_validation() -> ExecutionResult:
    return _execution(
        ["borough", "avg_fare", "orders"],
        [
            {"borough": "Manhattan", "avg_fare": 19.5, "orders": 102938},
            {"borough": "Queens", "avg_fare": 15.2, "orders": 88211},
        ],
    )


def test_validate_spec_accepts_bar_line_and_clamps_top_n(cfg: Nl2DataConfig) -> None:
    """bar/line with >=1 measure pass; top_n is clamped to the config cap."""
    execution = _exec_for_validation()
    spec, error = validate_spec(_BAR_SPEC, execution, cfg)
    assert error is None and spec is not None
    assert (spec.chart_type, spec.dimension, spec.measures) == ("bar", "borough", ("avg_fare",))
    line, error = validate_spec(
        {"chart_type": "line", "dimension": "borough", "measures": ["avg_fare", "orders"],
         "title": "t", "top_n": 999},
        execution, cfg,
    )
    assert error is None
    assert line.top_n == cfg.export.top_n_default  # clamped
    assert set(line.measures) == {"avg_fare", "orders"}


def test_validate_spec_rejections(cfg: Nl2DataConfig) -> None:
    """Every validator rule rejects with a readable reason."""
    execution = _exec_for_validation()
    cases = [
        ({"chart_type": "scatter", "dimension": "borough", "measures": ["avg_fare"]},
         "v1 仅 bar/line/pie"),
        ({"chart_type": "none"}, "不适合画图"),
        ({"chart_type": "bar", "dimension": "ghost", "measures": ["avg_fare"]}, "维度列不存在"),
        ({"chart_type": "bar", "dimension": "orders", "measures": ["avg_fare"]}, "类别/时间型"),
        ({"chart_type": "bar", "dimension": "borough", "measures": ["borough"]}, "数值型"),
        ({"chart_type": "bar", "dimension": "borough", "measures": []}, "measures 为空"),
        ({"chart_type": "pie", "dimension": "borough", "measures": ["avg_fare", "orders"]},
         "恰好 1 个度量"),
    ]
    for candidate, fragment in cases:
        spec, error = validate_spec(candidate, execution, cfg)
        assert spec is None, candidate
        assert fragment in (error or ""), (candidate, error)
    spec, error = validate_spec("not-a-dict", execution, cfg)
    assert spec is None and "JSON 对象" in (error or "")


def test_decide_chart_spec_disabled_makes_no_llm_call(
    cfg: Nl2DataConfig, monkeypatch: pytest.MonkeyPatch
) -> None:
    """chart_llm=False degrades to plain data export with zero LLM calls."""
    calls = _install_chat(monkeypatch, _result(_BAR_SPEC))
    spec, error = decide_chart_spec("q", _exec_for_validation(),
                                    replace(cfg, export=replace(cfg.export, chart_llm=False)))
    assert spec is None and "chart_llm 已关闭" in (error or "")
    assert calls == []


def test_decide_chart_spec_stub_roundtrip(
    cfg: Nl2DataConfig, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A valid LLM proposal passes through the validator; prompt carries schema."""
    calls = _install_chat(monkeypatch, _result(_BAR_SPEC))
    spec, error = decide_chart_spec("各行政区平均车费", _exec_for_validation(), cfg)
    assert error is None and spec == ChartSpec("bar", "borough", ("avg_fare",),
                                               "各行政区平均车费", 15)
    user_content = calls[0]["messages"][1]["content"]
    assert "borough(categorical)" in user_content and "avg_fare(numeric)" in user_content
    assert calls[0]["json_schema"] is not None


def test_decide_chart_spec_degradation_paths(
    cfg: Nl2DataConfig, monkeypatch: pytest.MonkeyPatch
) -> None:
    """LlmError / unparseable / invalid-proposal all degrade with reasons."""

    def _boom(*args: object, **kwargs: object) -> object:
        raise LlmError("api down", "api")

    monkeypatch.setattr("export.chart_spec.chat", _boom)
    spec, error = decide_chart_spec("q", _exec_for_validation(), cfg)
    assert spec is None and "LLM 判定调用失败(api)" in (error or "")

    _install_chat(monkeypatch, _result(None, content="not json"))
    spec, error = decide_chart_spec("q", _exec_for_validation(), cfg)
    assert spec is None and "无法按 JSON 解析" in (error or "")

    _install_chat(monkeypatch, _result({"chart_type": "pie", "dimension": "borough",
                                        "measures": ["avg_fare", "orders"]}))
    spec, error = decide_chart_spec("q", _exec_for_validation(), cfg)
    assert spec is None and "恰好 1 个度量" in (error or "")


# ------------------------------------------------------------ export run ----
def test_run_export_bundle_with_chart(cfg: Nl2DataConfig, monkeypatch: pytest.MonkeyPatch) -> None:
    """Full bundle: xlsx (data/meta/chart sheets + native chart) + manifest v1."""
    _install_chat(monkeypatch, _result(_BAR_SPEC))
    outcome = _outcome(
        _execution(
            ["borough", "avg_fare"],
            [
                {"borough": "Manhattan", "avg_fare": 19.5},
                {"borough": "Queens", "avg_fare": 15.2},
                {"borough": "Bronx", "avg_fare": 13.9},
            ],
        ),
        sql="SELECT borough, avg_fare FROM t",
    )
    result = run_export("各行政区平均车费", outcome, cfg)
    assert result.chart_spec is not None and result.chart_error is None
    assert not result.truncated and result.exported_rows == 3

    wb = load_workbook(result.xlsx_path)
    assert wb.sheetnames == ["data", "meta", "chart"]
    ws = wb["data"]
    assert [c.value for c in ws[1]] == ["borough", "avg_fare"]
    assert ws["A2"].value == "Manhattan" and ws.cell(row=4, column=2).value == 13.9
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:B4"
    charts = wb["chart"]._charts  # noqa: SLF001 -- openpyxl has no public accessor
    assert len(charts) == 1 and charts[0].title is not None

    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["version"] == 1
    assert manifest["chart_spec"]["chart_type"] == "bar"
    assert manifest["truncated"] is False and manifest["exported_rows"] == 3
    assert manifest["columns"] == [
        {"name": "borough", "dtype": "str", "role": "categorical"},
        {"name": "avg_fare", "dtype": "float", "role": "numeric"},
    ]
    for key in ("question", "sql", "source_tables", "tool_version", "generated_at"):
        assert manifest[key]
    assert Path(manifest["artifacts"]["xlsx"]).is_file()
    assert Path(manifest["artifacts"]["manifest"]).is_file()


def test_run_export_truncation_disclosure(
    cfg: Nl2DataConfig, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Mandatory visibility: truncated exports carry the meta warning row."""
    _install_chat(monkeypatch, _result(_BAR_SPEC))
    small = replace(cfg, export=replace(cfg.export, max_rows=2))
    outcome = _outcome(
        _execution(
            ["borough", "avg_fare"],
            [
                {"borough": "Manhattan", "avg_fare": 19.5},
                {"borough": "Queens", "avg_fare": 15.2},
                {"borough": "Bronx", "avg_fare": 13.9},
            ],
            rowcount=100,
        )
    )
    result = run_export("q", outcome, small)
    assert result.truncated and result.exported_rows == 2 and result.total_rows == 100

    wb = load_workbook(result.xlsx_path)
    # The chart still draws on the exported slice; truncation is disclosed.
    assert wb.sheetnames == ["data", "meta", "chart"]
    meta_keys = [wb["meta"].cell(row=i, column=1).value for i in range(1, 15)]
    assert "⚠ 已截断" in meta_keys
    warn_row = meta_keys.index("⚠ 已截断") + 1
    warn_value = wb["meta"].cell(row=warn_row, column=2).value
    assert "导出前 2 行" in warn_value and "共 100 行" in warn_value

    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["truncated"] is True
    assert (manifest["total_rows"], manifest["exported_rows"]) == (100, 2)


def test_run_export_reads_full_spill(cfg: Nl2DataConfig, monkeypatch: pytest.MonkeyPatch) -> None:
    """detail_ref Parquet is the preferred source: full rows over inline sample."""
    _install_chat(monkeypatch, _result(_BAR_SPEC))
    scratch = cfg.paths.scratch_dir
    scratch.mkdir(parents=True, exist_ok=True)
    spill = scratch / "spill.parquet"
    pd.DataFrame(
        {"borough": [f"z{i}" for i in range(5)], "avg_fare": [1.0, 2.0, 3.0, 4.0, 5.0]}
    ).to_parquet(spill)
    execution = _execution(
        ["borough", "avg_fare"],
        [{"borough": "z0", "avg_fare": 1.0}],
        rowcount=5,
        detail_ref="data/scratch/spill.parquet",
    )
    result = run_export("q", _outcome(execution), cfg)
    wb = load_workbook(result.xlsx_path)
    assert wb["data"].max_row == 6  # header + 5 full rows from the spill
    assert result.exported_rows == 5 and not result.truncated


def test_run_export_requires_success(cfg: Nl2DataConfig) -> None:
    """Exporting a non-executed outcome is a loud ValueError."""
    with pytest.raises(ValueError, match="成功执行"):
        run_export("q", QaOutcome(question="q"), cfg)


def test_chart_error_scrubs_env_values(
    cfg: Nl2DataConfig, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A canary value inside an LLM error never reaches disk (manifest/xlsx)."""
    canary = "sk-canary-deadbeef-not-a-real-key"
    monkeypatch.setenv("LLM_API_KEY", canary)

    def _boom(*args: object, **kwargs: object) -> object:
        raise LlmError(f"上游失败 upstream said {canary}", "api")

    monkeypatch.setattr("export.chart_spec.chat", _boom)
    outcome = _outcome(
        _execution(["borough", "avg_fare"], [{"borough": "a", "avg_fare": 1.0}])
    )
    result = run_export("q", outcome, cfg)
    assert result.chart_spec is None
    # The scrub replaced the canary with *** inside the persisted reason.
    assert result.chart_error is not None
    assert canary not in result.chart_error and "upstream said ***" in result.chart_error
    manifest_text = result.manifest_path.read_text(encoding="utf-8")
    assert canary not in manifest_text
    wb = load_workbook(result.xlsx_path)
    for row in wb["meta"].iter_rows(values_only=True):
        assert canary not in str(row)


# ------------------------------------------------------------------- CLI ----
@pytest.fixture(scope="module")
def cli_env(tmp_path_factory: pytest.TempPathFactory) -> Nl2DataConfig:
    """Workspace with one ingested table (profiles + cards) for CLI tests."""
    from catalog.cards import build_cards
    from catalog.profiler import run_profiles
    from ingest.common import PreparedTable, ingest_tables

    root = tmp_path_factory.mktemp("exportcli")
    config_file = root / "config.yaml"
    config_file.write_text(FIXTURE_YAML, encoding="utf-8")
    loaded = _load(config_file)
    ingest_tables(
        source_name="shop",
        source_type="excel",
        source_path=root / "shop.xlsx",
        tables=[
            PreparedTable(
                original_name="订单",
                frame=pd.DataFrame(
                    {
                        "订单ID": [1, 2, 3, 4],
                        "金额": [10.0, 20.0, 30.0, 40.0],
                        "城市": ["北京", "上海", "北京", "上海"],
                    }
                ),
            ),
        ],
        cfg=loaded,
    )
    run_profiles(loaded)
    build_cards(loaded)
    return loaded


def test_ask_single_shot_export_flag(
    cli_env: Nl2DataConfig, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """`ask --export xlsx` answers, then writes the bundle and prints paths."""
    from sqlgen.generate import SQLGeneration

    ok = SQLGeneration(
        sql="SELECT cheng_shi AS borough, sum(jin_e) AS total FROM ding_dan GROUP BY cheng_shi",
        needs_clarification=False, clarification=None, candidates_tried=1,
    )
    monkeypatch.setattr("nl2data.qa.generate", lambda *a, **k: ok)
    spec = {
        "chart_type": "bar", "dimension": "borough", "measures": ["total"],
        "title": "城市金额", "top_n": 5,
    }
    _install_chat(monkeypatch, _result(spec))
    config_file = cli_env.paths.data_dir.parent / "config.yaml"
    result = runner.invoke(
        app,
        ["ask", "按城市汇总金额", "--no-interpret", "--export", "xlsx", "--config",
         str(config_file)],
    )
    assert result.exit_code == 0, result.output
    assert "已导出" in result.output and "清单" in result.output
    bundles = list(cli_env.export.dir.glob("*/result.xlsx"))
    assert len(bundles) == 1 and bundles[0].with_name("manifest.json").is_file()


def test_ask_export_rejects_unknown_format(cli_env: Nl2DataConfig) -> None:
    """--export pdf fails fast with exit code 2."""
    config_file = cli_env.paths.data_dir.parent / "config.yaml"
    result = runner.invoke(app, ["ask", "q", "--export", "pdf", "--config", str(config_file)])
    assert result.exit_code == 2
    assert "不支持的导出格式" in result.output


def test_export_cli_dispatch_csv_and_xlsx(
    cli_env: Nl2DataConfig, monkeypatch: pytest.MonkeyPatch
) -> None:
    """/export dispatch: csv requires a path; xlsx needs a last outcome."""
    from nl2data.cli import _export_cli

    captured: list[str] = []

    def _capture(msg: object) -> None:
        captured.append(str(msg))

    monkeypatch.setattr("nl2data.cli.console.print", _capture)
    _export_cli("/export", cli_env, None, "")  # no format → usage hint
    assert any("用法" in line for line in captured)
    captured.clear()
    _export_cli("/export csv", cli_env, None, "")  # csv without path → usage hint
    assert any("用法" in line for line in captured)
    captured.clear()
    _export_cli("/export xlsx", cli_env, None, "")  # nothing to export yet
    assert any("没有可导出的成功结果" in line for line in captured)


# --------------------------------------------------------------- real run ----
@pytest.mark.slow
def test_export_real_llm_chart_bundle() -> None:
    """Real LLM chart decision over the repo warehouse (manual inspection).

    Skips without real credentials; run with a valid LLM key to produce one
    bundle for human spot-checking of the native chart rendering.
    """
    import os

    from catalog.cards import build_cards
    from catalog.profiler import run_profiles
    from ingest.parquet import ingest_parquet
    from nl2data.qa import ask_once

    for key in ("LLM_BASE_URL", "LLM_API_KEY", "LLM_MODEL"):
        if not os.environ.get(key, "").strip():
            pytest.skip("real LLM_* environment not set")

    repo_cfg = _load(Path(__file__).resolve().parent.parent / "config.yaml")
    if not repo_cfg.paths.warehouse.is_file():
        ingest_parquet(
            Path(__file__).resolve().parent.parent / "samples" / "nyc-taxi"
            / "yellow_tripdata.parquet",
            repo_cfg,
            name="yellow_tripdata",
        )
        run_profiles(repo_cfg)
        build_cards(repo_cfg)
    question = "各个行政区的黄车订单量是多少?"
    outcome = ask_once(question, repo_cfg, no_interpret=True)
    assert outcome.ok, outcome.failure_reason
    result = run_export(question, outcome, repo_cfg)
    print(f"xlsx={result.xlsx_path} chart={result.chart_spec} err={result.chart_error}")
    assert result.xlsx_path.is_file() and result.manifest_path.is_file()
