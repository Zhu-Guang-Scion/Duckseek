"""Excel (.xlsx/.xlsm) ingestion: sheet frames into the shared pipeline.

Reader strategy: the DuckDB ``excel`` extension is preferred (fast, typed
reads); any failure falls back to ``pandas.read_excel`` (openpyxl engine) with
a warning. Header rows are auto-detected by scanning the first
``cfg.ingest.header_scan_rows`` rows with openpyxl (read-only), so report
titles and blank leading rows never become column names.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import duckdb
import pandas as pd
from openpyxl import load_workbook

from catalog.naming import clean_name
from catalog.store import SourceEntry
from ingest.common import SOURCE_TYPE_EXCEL, IngestError, PreparedTable, ingest_tables
from nl2data.config import Nl2DataConfig

logger = logging.getLogger(__name__)

_SUPPORTED_SUFFIXES = frozenset({".xlsx", ".xlsm"})
_EXCEL_EXTENSION = "excel"
_READ_XLSX_SQL = "SELECT * FROM read_xlsx(?, sheet = ?, header = true, stop_at_empty = false)"

# None until the first attempt; True/False afterwards (never retried).
_extension_available: bool | None = None


def detect_header_row(rows: list[list[Any]], *, max_scan: int = 5) -> int:
    """Return the 0-based index of the header row inside ``rows``.

    A row qualifies as the header when it holds at least two non-empty cells;
    blank rows and single-cell "report title" rows are skipped. When no row
    qualifies within ``max_scan`` rows, ``max_scan`` is returned (scan cap).

    Args:
        rows: Leading sheet rows as raw cell values.
        max_scan: Maximum number of rows to inspect.

    Returns:
        The 0-based header row index, or ``max_scan`` when capped.
    """
    for index, row in enumerate(rows[:max_scan]):
        non_empty = sum(1 for value in row if value is not None and str(value).strip() != "")
        if non_empty >= 2:
            return index
    return max_scan


def _load_excel_extension() -> bool:
    """INSTALL and LOAD the DuckDB excel extension once per process.

    Returns:
        True when the extension is usable. Any failure is logged once as a
        warning and never retried; the pandas fallback then serves sheets.
    """
    global _extension_available
    if _extension_available is not None:
        return _extension_available
    try:
        conn = duckdb.connect()
        try:
            conn.execute(f"INSTALL {_EXCEL_EXTENSION}")
            conn.execute(f"LOAD {_EXCEL_EXTENSION}")
        finally:
            conn.close()
    except duckdb.Error as exc:
        logger.warning("duckdb excel extension unavailable; using pandas fallback: %s", exc)
        _extension_available = False
    else:
        _extension_available = True
    return _extension_available


def _sheet_names(source: Path) -> list[str]:
    """Return the workbook sheet names via a cheap read-only openpyxl pass."""
    workbook = load_workbook(source, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _probe_header_row(source: Path, sheet: str, scan_rows: int) -> int:
    """Scan the first ``scan_rows`` rows of ``sheet`` for the header row."""
    workbook = load_workbook(source, read_only=True)
    try:
        worksheet = workbook[sheet]
        rows = [
            list(row)
            for row in worksheet.iter_rows(min_row=1, max_row=scan_rows, values_only=True)
        ]
    finally:
        workbook.close()
    return detect_header_row(rows, max_scan=scan_rows)


def _read_sheet_via_duckdb(source: Path, sheet: str) -> pd.DataFrame:
    """Read one sheet with the DuckDB excel extension (header on row 1)."""
    conn = duckdb.connect()
    try:
        conn.execute(f"LOAD {_EXCEL_EXTENSION}")
        return conn.sql(_READ_XLSX_SQL, params=[str(source), sheet]).df()
    finally:
        conn.close()


def _is_blank_header(name: object) -> bool:
    """Return True when a column header cell is empty or a pandas placeholder."""
    if name is None or (isinstance(name, float) and pd.isna(name)):
        return True
    text = str(name).strip()
    return text == "" or text.lower().startswith("unnamed")


def _postprocess(frame: pd.DataFrame, *, sheet: str) -> pd.DataFrame:
    """Shared cleanup for both reader paths.

    Drops columns whose header is blank and whose data is entirely empty,
    drops fully empty data rows, and downgrades header-bearing all-NULL
    columns to ``string`` (logged) so they survive the Parquet round-trip.

    Args:
        frame: Raw sheet frame from either reader.
        sheet: Sheet name, for log messages.

    Returns:
        The cleaned frame.
    """
    if frame.empty:
        return frame
    blank = [
        str(col)
        for col in frame.columns
        if _is_blank_header(col) and frame[col].isna().all()
    ]
    if blank:
        frame = frame.drop(columns=blank)
    frame = frame.dropna(how="all").reset_index(drop=True)
    for column in frame.columns:
        if frame[column].isna().all():
            frame[column] = frame[column].astype("string")
            logger.info(
                "column %s in sheet %s has no non-NULL values; downgraded to VARCHAR",
                column,
                sheet,
            )
    return frame


def _read_sheet(
    source: Path,
    sheet: str,
    cfg: Nl2DataConfig,
    extension_available: bool,
) -> pd.DataFrame | None:
    """Read and clean one sheet; return None when it ends up empty.

    The DuckDB reader is only tried when the detected header row is the first
    row (its ``header=true`` contract); any ``duckdb.Error`` such as the
    Binder Error raised for empty sheets falls back to pandas with a warning.
    """
    header_row = _probe_header_row(source, sheet, cfg.ingest.header_scan_rows)
    frame: pd.DataFrame | None = None
    if extension_available and header_row == 0:
        try:
            frame = _read_sheet_via_duckdb(source, sheet)
        except duckdb.Error as exc:
            logger.warning(
                "duckdb read_xlsx failed for sheet %r in %s; falling back to pandas: %s",
                sheet,
                source,
                exc,
            )
    if frame is None:
        frame = pd.read_excel(source, sheet_name=sheet, engine="openpyxl", header=header_row)
    frame = _postprocess(frame, sheet=sheet)
    if frame.empty:
        logger.warning("sheet %r in %s is empty after cleanup; skipped", sheet, source)
        return None
    return frame


def ingest_excel(
    source: Path,
    cfg: Nl2DataConfig,
    sheet: str | None = None,
    name: str | None = None,
) -> SourceEntry:
    """Ingest an ``.xlsx``/``.xlsm`` workbook through the shared pipeline.

    Every non-empty sheet becomes one table: the header row is auto-detected,
    the sheet is read via the DuckDB excel extension when available (pandas
    fallback otherwise), then handed to :func:`ingest.common.ingest_tables`,
    which owns naming, Parquet storage, DuckDB views and catalog lineage.

    Args:
        source: Workbook path; suffix must be ``.xlsx`` or ``.xlsm``.
        cfg: Active configuration.
        sheet: Exact sheet name to ingest; ``None`` ingests every sheet.
        name: Optional source alias overriding the file stem for the slug.

    Returns:
        The catalog entry written for this source.

    Raises:
        IngestError: If the file is missing, its suffix is unsupported
            (``.xls``/``.xlsb``/``.numbers`` are not supported), the requested
            sheet does not exist, or no non-empty sheet remains.
    """
    source = Path(source)
    if not source.is_file():
        msg = f"Excel file not found: {source}"
        raise IngestError(msg)
    suffix = source.suffix.lower()
    if suffix not in _SUPPORTED_SUFFIXES:
        msg = (
            f"unsupported Excel file {source.name!r} (suffix {suffix!r}); "
            "supported: .xlsx/.xlsm (.xls/.xlsb/.numbers are not supported)"
        )
        raise IngestError(msg)

    available = _sheet_names(source)
    if sheet is not None and sheet not in available:
        msg = (
            f"sheet {sheet!r} not found in {source}; "
            f"available sheets: {', '.join(repr(n) for n in available)}"
        )
        raise IngestError(msg)
    wanted = available if sheet is None else [sheet]

    extension_available = _load_excel_extension()
    tables: list[PreparedTable] = []
    for sheet_name in wanted:
        frame = _read_sheet(source, sheet_name, cfg, extension_available)
        if frame is not None:
            tables.append(PreparedTable(original_name=sheet_name, frame=frame))

    return ingest_tables(
        source_name=clean_name(
            name or source.stem,
            max_length=cfg.ingest.max_name_length,
        ),
        source_type=SOURCE_TYPE_EXCEL,
        source_path=source,
        tables=tables,
        cfg=cfg,
    )
